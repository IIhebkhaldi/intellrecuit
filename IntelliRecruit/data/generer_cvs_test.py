# -*- coding: utf-8 -*-
"""
Génère le dataset de test annoté (S6/S7) : data/sample_cvs/ + data/ground_truth.json.

⭐ Piloté par les données : chaque profil produit À LA FOIS le fichier CV et son
entrée de vérité terrain. L'annotation est donc toujours cohérente avec le CV.
Élargi à 100 CVs (exigence cahier §5.1 : 100+ CVs annotés ; 6 -> 30 -> 70 -> 100).

Lancer :  python data/generer_cvs_test.py
(nécessite fpdf2 et python-docx)
"""
import io
import json
import sys
from pathlib import Path

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8")

DOSSIER = Path(__file__).parent / "sample_cvs"
DOSSIER.mkdir(exist_ok=True)
GROUND_TRUTH = Path(__file__).parent / "ground_truth.json"


# ------------------------------------------------------------------
# Métiers -> compétences (toutes présentes dans l'ontologie skills_normalizer)
# ------------------------------------------------------------------
ROLES = {
    "data_scientist": dict(
        titre_fr="Data Scientist", titre_en="Data Scientist",
        skills=["python", "sql", "machine learning", "pandas", "scikit-learn"],
        field_fr="Data & IA", field_en="Data Science",
        societes=["Vermeg", "Sofrecom", "Talan", "Sopra HR"],
    ),
    "backend": dict(
        titre_fr="Développeur Backend", titre_en="Backend Developer",
        skills=["python", "django", "fastapi", "postgresql", "docker"],
        field_fr="Génie Logiciel", field_en="Software Engineering",
        societes=["FinTech Ltd", "Instadeep", "Expensya", "Vneuron"],
    ),
    "frontend": dict(
        titre_fr="Développeur Frontend", titre_en="Frontend Developer",
        skills=["javascript", "react", "typescript", "angular", "node.js"],
        field_fr="Développement Web", field_en="Web Development",
        societes=["Wevioo", "Telnet", "Focus Corporation"],
    ),
    "bi": dict(
        titre_fr="Analyste Business Intelligence", titre_en="BI Analyst",
        skills=["power bi", "sql server", "ssis", "dax", "excel", "talend"],
        field_fr="Business Intelligence", field_en="Business Intelligence",
        societes=["Sofrecom", "BFI Group", "Talan"],
    ),
    "ml": dict(
        titre_fr="Ingénieur Machine Learning", titre_en="Machine Learning Engineer",
        skills=["python", "tensorflow", "keras", "nlp", "kubernetes", "mlops"],
        field_fr="Intelligence Artificielle", field_en="Artificial Intelligence",
        societes=["DataCorp", "Instadeep", "Clevy"],
    ),
    "devops": dict(
        titre_fr="Ingénieur DevOps", titre_en="DevOps Engineer",
        skills=["docker", "kubernetes", "terraform", "aws", "jenkins", "linux"],
        field_fr="Cloud & Infrastructure", field_en="Cloud & Infrastructure",
        societes=["Orange", "Ooredoo", "AWS Partner"],
    ),
    "fullstack": dict(
        titre_fr="Développeur Full Stack", titre_en="Full Stack Developer",
        skills=["javascript", "react", "node.js", "mongodb", "docker"],
        field_fr="Développement Web", field_en="Web Development",
        societes=["TunisAI", "Enova Robotics", "GoMyCode"],
    ),
    "data_engineer": dict(
        titre_fr="Data Engineer", titre_en="Data Engineer",
        skills=["python", "spark", "airflow", "sql", "kafka", "hadoop"],
        field_fr="Data Engineering", field_en="Data Engineering",
        societes=["Talan", "Devoteam", "Capgemini"],
    ),
    "java": dict(
        titre_fr="Développeur Java", titre_en="Java Developer",
        skills=["java", "spring", "mysql", "docker", "git", "uml"],
        field_fr="Génie Logiciel", field_en="Software Engineering",
        societes=["Sopra Steria", "Accenture", "Proxym"],
    ),
    "mobile": dict(
        titre_fr="Développeur Mobile", titre_en="Mobile Developer",
        skills=["kotlin", "swift", "java", "git"],
        field_fr="Développement Mobile", field_en="Mobile Development",
        societes=["Digitup", "Konnect", "Kaoun"],
    ),
}

ECOLES_FR = ["ESPRIT", "INSAT", "ENIT", "ISET Sousse", "FST Tunis", "SUPCOM"]
ECOLES_EN = ["University of Manchester", "Sorbonne Université",
             "Technical University of Berlin", "Polytechnique Montréal"]

VILLES = [
    ("Tunis", "Tunisie"), ("Ariana", "Tunisie"), ("Sousse", "Tunisie"),
    ("Sfax", "Tunisie"), ("Paris", "France"), ("Lyon", "France"),
    ("London", "UK"), ("Berlin", "Allemagne"), ("Madrid", "Espagne"),
    ("Montréal", "Canada"), ("Milan", "Italie"), ("Casablanca", "Maroc"),
]

NOMS = [
    ("Nour", "Belhadj"), ("Aymen", "Gharbi"), ("Rania", "Mansour"),
    ("Skander", "Bouazizi"), ("Ines", "Khemiri"), ("Bilel", "Jaziri"),
    ("Farah", "Ncibi"), ("Oussama", "Zouari"), ("Maryem", "Chaabane"),
    ("Wassim", "Trabelsi"), ("Emma", "Laurent"), ("Lucas", "Martin"),
    ("Sofia", "Rossi"), ("David", "Meyer"), ("Carla", "Garcia"),
    ("Hedi", "Slama"), ("Nadia", "Ben Youssef"), ("Karim", "Haddad"),
    ("Sami", "Ferjani"), ("Yasmine", "Ayari"), ("Mehdi", "Bouzid"),
    ("Leila", "Marzouki"), ("Anis", "Guesmi"), ("Dorra", "Hamdi"),
    # --- élargissement à 70 CVs (retour encadrant : viser 100+) ---
    ("Ahmed", "Khelifi"), ("Salma", "Riahi"), ("Firas", "Nasri"),
    ("Amine", "Chihi"), ("Rim", "Baccouche"), ("Hamza", "Dridi"),
    ("Ghada", "Louati"), ("Marwa", "Sassi"), ("Zied", "Fakhfakh"),
    ("Nesrine", "Chebbi"), ("Walid", "Amri"), ("Chaima", "Belaid"),
    ("Hatem", "Karray"), ("Sarra", "Mejri"), ("Bassem", "Tlili"),
    ("Olfa", "Rekik"), ("Ramzi", "Bouaziz"), ("Amani", "Jlassi"),
    ("Seif", "Ghariani"), ("Mariem", "Ouali"), ("Tarek", "Benslimane"),
    ("Asma", "Kallel"), ("Nizar", "Toumi"), ("Hiba", "Fourati"),
    ("Malek", "Ayadi"), ("Julien", "Dubois"), ("Chloe", "Bernard"),
    ("Marco", "Bianchi"), ("Elena", "Popescu"), ("Thomas", "Weber"),
    ("Laura", "Sanchez"), ("Antoine", "Moreau"), ("Giulia", "Conti"),
    ("Miguel", "Torres"), ("Sophie", "Muller"), ("Paul", "Lefevre"),
    ("Ana", "Silva"), ("Nicolas", "Girard"), ("Marta", "Ferrari"),
    ("Kevin", "Schmidt"),
    # --- élargissement à 100 CVs (exigence cahier : 100+ CVs annotés) ---
    ("Mouna", "Jebali"), ("Slim", "Khemakhem"), ("Rania", "Zaidi"),
    ("Fedi", "Chaabene"), ("Syrine", "Larbi"), ("Aziz", "Mabrouk"),
    ("Emna", "Trigui"), ("Rayen", "Souissi"), ("Cyrine", "Nasr"),
    ("Mootez", "Brahmi"), ("Nour", "Dhaou"), ("Ayoub", "Sellami"),
    ("Ranim", "Haj Ali"), ("Zeineb", "Missaoui"), ("Oumaima", "Chaari"),
    ("Youssra", "Belaifa"), ("Aziz", "Ncib"), ("Wael", "Boukadida"),
    ("Ilef", "Gharsallah"), ("Bacem", "Sfar"), ("Rihab", "Chouchane"),
    ("Achraf", "Zribi"), ("Ismail", "Ouertani"), ("Sirine", "Abbes"),
    ("Hamdi", "Rahmani"), ("Lucia", "Romano"), ("Diego", "Fernandez"),
    ("Hannah", "Fischer"), ("Julia", "Kowalski"), ("Pedro", "Costa"),
]


def _lien_linkedin(prenom, nom):
    return f"{prenom}-{nom}".lower().replace(" ", "-")


# ------------------------------------------------------------------
# Les 6 CVs "historiques" (gardés à l'identique) + génération des autres
# ------------------------------------------------------------------
PROFILS = [
    dict(filename="cv_amira_bensalem.pdf", format="pdf", language="fr",
         first_name="Amira", last_name="Ben Salem",
         email="amira.bensalem@gmail.com", phone="+216 22 345 678",
         location="Ariana, Tunisie", linkedin="amira-bensalem",
         titre="Data Scientist Junior",
         skills=["python", "sql", "machine learning", "power bi", "pandas", "tensorflow"],
         languages=[("arabic", "natif"), ("french", "courant"), ("english", "professionnel")],
         experience=[("Stage Data Analyst", "Vermeg", "Juin 2025", "Août 2025",
                      "Analyse des données clients avec Python et Power BI.")],
         education=[("Ingénieur Informatique", "Data & IA", "ESPRIT", 2022, 2025)],
         certifications=["Microsoft Power BI Data Analyst (PL-300)",
                         "TensorFlow Developer Certificate"]),
    dict(filename="cv_john_smith.pdf", format="pdf", language="en",
         first_name="John", last_name="Smith",
         email="john.smith@outlook.com", phone="+44 20 7946 0958",
         location="London, UK", linkedin="john-smith-dev",
         titre="Senior Backend Developer",
         skills=["python", "django", "postgresql", "docker", "aws"],
         languages=[("english", "native"), ("french", "intermediate")],
         experience=[("Backend Developer", "FinTech Ltd", "2021", "2024",
                      "Built and maintained REST APIs serving millions of requests.")],
         education=[("BSc Computer Science", None, "University of Manchester", 2017, 2020)],
         certifications=["AWS Certified Solutions Architect"]),
    dict(filename="cv_lina_kacem.pdf", format="pdf", language="fr",
         first_name="Lina", last_name="Kacem",
         email="lina.kacem@etudiant.esprit.tn", phone=None,
         location=None, linkedin=None,
         titre="Étudiante en Data Science",
         skills=["python", "pandas", "scikit-learn"],
         languages=[("arabic", None), ("french", None), ("english", None)],
         experience=[],
         education=[("Ingénieur Informatique", None, "ESPRIT", 2023, 2026)],
         certifications=[]),
    dict(filename="cv_mohamed_trabelsi.docx", format="docx", language="fr",
         first_name="Mohamed", last_name="Trabelsi",
         email="mohamed.trabelsi@yahoo.fr", phone="+216 98 765 432",
         location="Sousse, Tunisie", linkedin=None,
         titre="Développeur Full Stack",
         skills=["javascript", "react", "node.js", "mongodb", "docker"],
         languages=[("arabic", "natif"), ("french", "courant"), ("english", "intermédiaire")],
         experience=[("Développeur Web", "TunisAI", "Janvier 2023", None,
                      "Développement d'une plateforme SaaS en React et Node.js.")],
         education=[("Licence en Informatique", None, "ISET Sousse", 2019, 2022)],
         certifications=[]),
    dict(filename="cv_sara_gharbi.docx", format="docx", language="en",
         first_name="Sara", last_name="Gharbi",
         email="sara.gharbi@protonmail.com", phone="+33 6 12 34 56 78",
         location="Paris, France", linkedin="sara-gharbi-ml",
         titre="Machine Learning Engineer",
         skills=["python", "tensorflow", "keras", "nlp", "mlops", "kubernetes", "machine learning"],
         languages=[("french", "native"), ("arabic", "native"), ("english", "fluent")],
         experience=[("ML Engineer", "DataCorp", "2022", None,
                      "Deployed NLP models for document classification in production.")],
         education=[("MSc Artificial Intelligence", None, "Sorbonne Université", 2020, 2022)],
         certifications=[]),
    dict(filename="cv_youssef_benali.txt", format="txt", language="fr",
         first_name="Youssef", last_name="Ben Ali",
         email="youssef.benali@hotmail.com", phone="55 123 456",
         location="Tunis, Tunisie", linkedin=None,
         titre="Analyste Business Intelligence",
         skills=["power bi", "sql server", "ssis", "dax", "excel", "talend",
                 "etl", "business intelligence"],
         languages=[("arabic", "natif"), ("french", "courant"), ("english", "B2")],
         experience=[("Consultant BI", "Sofrecom", "Mars 2024", None,
                      "Conception de tableaux de bord et de pipelines ETL en business intelligence.")],
         education=[("Master Business Intelligence", None, "Université de Tunis", 2021, 2023)],
         certifications=[]),
]


def _profils_generes():
    """Compose 94 profils supplémentaires à partir des pools (annotation dérivée)."""
    profils = []
    roles = list(ROLES)
    formats = ["pdf", "docx", "txt"]
    domaines = ["gmail.com", "outlook.com", "yahoo.fr", "hotmail.com"]
    for i, (prenom, nom) in enumerate(NOMS):
        role = ROLES[roles[i % len(roles)]]
        langue = "fr" if i % 3 != 2 else "en"  # 2/3 fr, 1/3 en
        fmt = formats[i % 3]
        ville, pays = VILLES[i % len(VILLES)]
        ecole = (ECOLES_FR if langue == "fr" else ECOLES_EN)[i % 4]
        tel = ["+216 " + f"{20 + i} {100 + i:03d} {200 + i:03d}",
               "+33 6 12 34 56 78", None][i % 3]
        linkedin = _lien_linkedin(prenom, nom) if i % 4 != 3 else None
        annee_fin = 2020 + (i % 5)
        profils.append(dict(
            filename=f"cv_{prenom.lower()}_{nom.lower().replace(' ', '_')}.{fmt}",
            format=fmt, language=langue,
            first_name=prenom, last_name=nom,
            email=f"{prenom.lower()}.{nom.lower().replace(' ', '')}@{domaines[i % 4]}",
            phone=tel, location=f"{ville}, {pays}", linkedin=linkedin,
            titre=role["titre_fr"] if langue == "fr" else role["titre_en"],
            skills=list(role["skills"]),
            languages=([("arabic", "natif"), ("french", "courant"), ("english", "professionnel")]
                       if langue == "fr" else [("english", "native"), ("french", "intermediate")]),
            experience=[(
                role["titre_fr"] if langue == "fr" else role["titre_en"],
                role["societes"][i % len(role["societes"])],
                "2022" if langue == "en" else "Janvier 2022",
                str(annee_fin) if langue == "en" else None,
                ("Contribution à des projets d'équipe et livraisons régulières."
                 if langue == "fr" else "Delivered features and collaborated with the team."),
            )],
            education=[(
                "Ingénieur Informatique" if langue == "fr" else "MSc Computer Science",
                role["field_fr"] if langue == "fr" else role["field_en"],
                ecole, annee_fin - 3, annee_fin,
            )],
            certifications=[],
        ))
    return profils


PROFILS += _profils_generes()


# ------------------------------------------------------------------
# Construction du texte d'un CV à partir d'un profil
# ------------------------------------------------------------------
def construire_texte(p) -> str:
    fr = p["language"] == "fr"
    L = []
    L.append(f"{p['first_name'].upper()} {p['last_name'].upper()}")
    L.append(p["titre"])
    if p["location"]:
        L.append(p["location"])
    L.append(("Email : " if fr else "Email: ") + p["email"])
    if p["phone"]:
        L.append(("Tél : " if fr else "Phone: ") + p["phone"])
    if p["linkedin"]:
        L.append(f"linkedin.com/in/{p['linkedin']}")
    L.append("")

    L.append("COMPÉTENCES" if fr else "SKILLS")
    L.append(", ".join(s.title() if " " not in s else s for s in p["skills"]))
    L.append("")

    if p["experience"]:
        L.append("EXPÉRIENCE PROFESSIONNELLE" if fr else "WORK EXPERIENCE")
        for titre, societe, debut, fin, desc in p["experience"]:
            sep = " chez " if fr else " at "
            borne = fin if fin else ("présent" if fr else "present")
            L.append(f"{titre}{sep}{societe} ({debut} - {borne})")
            L.append(desc)
        L.append("")

    L.append("FORMATION" if fr else "EDUCATION")
    for degree, field, school, sy, ey in p["education"]:
        if field and fr:
            L.append(f"{degree}, spécialité {field} — {school} ({sy} - {ey})")
        elif fr:
            L.append(f"{degree} — {school} ({sy} - {ey})")
        else:
            L.append(f"{degree}, {school} ({sy} - {ey})")
    L.append("")

    L.append("LANGUES" if fr else "LANGUAGES")
    noms_fr = {"arabic": "Arabe", "french": "Français", "english": "Anglais"}
    noms_en = {"arabic": "Arabic", "french": "French", "english": "English"}
    noms = noms_fr if fr else noms_en
    items = []
    for nom, niveau in p["languages"]:
        items.append(f"{noms.get(nom, nom.title())}" + (f" ({niveau})" if niveau else ""))
    L.append(", ".join(items))

    if p["certifications"]:
        L.append("")
        L.append("CERTIFICATIONS")
        L.extend(p["certifications"])

    return "\n".join(L) + "\n"


# ------------------------------------------------------------------
# Écriture des fichiers
# ------------------------------------------------------------------
def creer_pdf(texte, chemin):
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=11)
    for ligne in texte.split("\n"):
        ligne = ligne.replace("—", "-").replace("’", "'")  # hors latin-1
        pdf.cell(0, 7, ligne, new_x="LMARGIN", new_y="NEXT")
    pdf.output(str(chemin))


def creer_docx(texte, chemin):
    from docx import Document
    doc = Document()
    for ligne in texte.split("\n"):
        doc.add_paragraph(ligne)
    doc.save(str(chemin))


def creer_txt(texte, chemin):
    chemin.write_text(texte, encoding="utf-8")


GENERATEURS = {"pdf": creer_pdf, "docx": creer_docx, "txt": creer_txt}


# ------------------------------------------------------------------
# Vérité terrain dérivée du profil
# ------------------------------------------------------------------
def entree_ground_truth(p) -> dict:
    return {
        "language": p["language"],
        "first_name": p["first_name"],
        "last_name": p["last_name"],
        "email": p["email"],
        "phone": p["phone"],
        "location": p["location"],
        "linkedin": f"https://www.linkedin.com/in/{p['linkedin']}" if p["linkedin"] else None,
        "skills": p["skills"],
        "languages": [nom for nom, _ in p["languages"]],
        "nb_experience": len(p["experience"]),
        "nb_education": len(p["education"]),
    }


if __name__ == "__main__":
    # nettoie les anciens CVs pour repartir propre
    for ancien in DOSSIER.glob("cv_*"):
        ancien.unlink()

    ground_truth = {
        "_description": "Vérité terrain annotée — générée depuis data/generer_cvs_test.py "
                        "(le CV et son annotation viennent de la même source, donc cohérents)."
    }
    for p in PROFILS:
        texte = construire_texte(p)
        GENERATEURS[p["format"]](texte, DOSSIER / p["filename"])
        ground_truth[p["filename"]] = entree_ground_truth(p)

    GROUND_TRUTH.write_text(
        json.dumps(ground_truth, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(f"{len(PROFILS)} CVs générés dans {DOSSIER}")
    print(f"Vérité terrain écrite dans {GROUND_TRUTH}")
