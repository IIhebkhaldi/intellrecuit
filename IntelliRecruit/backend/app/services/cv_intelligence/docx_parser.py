# -*- coding: utf-8 -*-
"""
Extraction du texte d'un fichier Word (.docx).

On lit les paragraphes ET les cellules de tableaux : beaucoup de CVs Word
sont mis en page à l'aide de tableaux (cahier des charges §5.1, A1).
"""
import io


def extraire_texte_docx(contenu: bytes) -> str:
    """DOCX (bytes) -> texte (paragraphes + cellules de tableaux)."""
    try:
        from docx import Document
        doc = Document(io.BytesIO(contenu))
    except Exception as e:
        raise ValueError(f"Impossible de lire le DOCX : {e}")

    lignes = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                lignes.append(cell.text)
    return "\n".join(l for l in lignes if l.strip())
