# -*- coding: utf-8 -*-
"""Tests S3 du service cv_intelligence (parsing multi-format + identité basique)."""
import io

import pytest

from app.services.cv_intelligence import (
    detecter_format,
    detecter_langue,
    extraire_email,
    extraire_identite,
    extraire_localisation,
    extraire_linkedin,
    extraire_nom_prenom,
    extraire_telephone,
    extraire_texte,
    parser_cv,
)

CV_TEXTE_FR = """AMIRA BEN SALEM
Data Scientist Junior
Ariana, Tunisie
Email : amira.bensalem@gmail.com
Tél : +216 22 345 678
linkedin.com/in/amira-bensalem

EXPÉRIENCE PROFESSIONNELLE
Stage Data Analyst chez Vermeg de Juin 2025 à Août 2025.
"""


# ---------- détection de format ----------
def test_detecter_format():
    assert detecter_format("cv.pdf") == "pdf"
    assert detecter_format("CV.DOCX") == "docx"
    assert detecter_format("notes.txt") == "txt"


def test_detecter_format_refuse():
    with pytest.raises(ValueError):
        detecter_format("photo.png")


# ---------- extraction de texte multi-format ----------
def test_extraire_texte_txt_utf8():
    assert "téléphone" in extraire_texte("a.txt", "téléphone".encode("utf-8"))


def test_extraire_texte_txt_cp1252():
    """Fichier créé sous Windows (Bloc-notes) : pas de crash, accents préservés."""
    assert "téléphone" in extraire_texte("a.txt", "téléphone".encode("cp1252"))


def test_extraire_texte_docx():
    from docx import Document
    doc = Document()
    doc.add_paragraph("MOHAMED TRABELSI")
    doc.add_paragraph("mohamed@yahoo.fr")
    buf = io.BytesIO()
    doc.save(buf)
    texte = extraire_texte("cv.docx", buf.getvalue())
    assert "MOHAMED TRABELSI" in texte
    assert "mohamed@yahoo.fr" in texte


def test_extraire_texte_docx_tableau():
    """Beaucoup de CVs Word sont mis en page avec des tableaux."""
    from docx import Document
    doc = Document()
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "SARA GHARBI"
    table.rows[0].cells[1].text = "+33 6 12 34 56 78"
    buf = io.BytesIO()
    doc.save(buf)
    texte = extraire_texte("cv.docx", buf.getvalue())
    assert "SARA GHARBI" in texte
    assert "+33 6 12 34 56 78" in texte


def test_extraire_texte_docx_corrompu():
    with pytest.raises(ValueError):
        extraire_texte("cv.docx", b"pas un vrai docx")


def test_pdf_hyperlien_cliquable_recupere():
    """Un lien LinkedIn CLIQUABLE (URL absente du texte visible) doit être
    récupéré depuis l'annotation du PDF, puis extrait comme profil."""
    from fpdf import FPDF

    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=12)
    pdf.cell(0, 8, "MARIE DURAND", new_x="LMARGIN", new_y="NEXT")
    pdf.cell(0, 8, "marie.durand@gmail.com", new_x="LMARGIN", new_y="NEXT")
    # le mot "LinkedIn" est cliquable ; l'URL n'est PAS écrite dans le texte
    pdf.cell(0, 8, "LinkedIn",
             link="https://www.linkedin.com/in/marie-durand",
             new_x="LMARGIN", new_y="NEXT")
    contenu = bytes(pdf.output())

    cv = parser_cv("marie.pdf", contenu)
    assert cv.identity.linkedin == "https://www.linkedin.com/in/marie-durand"


# ---------- identité : téléphone ----------
def test_telephone_tunisien():
    assert extraire_telephone("Tél : +216 22 345 678") == "+216 22 345 678"


def test_telephone_francais_groupes_courts():
    """Bug corrigé en démo : les groupes à 1 chiffre (+33 6 ...) sont gérés."""
    assert extraire_telephone("Phone: +33 6 12 34 56 78") == "+33 6 12 34 56 78"


def test_telephone_ignore_annees():
    """Une plage d'années ne doit PAS être prise pour un téléphone."""
    assert extraire_telephone("Projet mené de 2021 - 2024 puis 2025") is None


def test_telephone_absent():
    assert extraire_telephone("Aucun numéro ici.") is None


# ---------- identité : autres champs ----------
def test_extraire_email_premier():
    assert extraire_email("a@b.tn puis c@d.fr") == "a@b.tn"


def test_extraire_nom_prenom_convention():
    """Convention J10 : 1er mot = prénom, le reste = nom (composé inclus)."""
    prenom, nom = extraire_nom_prenom(CV_TEXTE_FR)
    assert prenom == "Amira"
    assert nom == "Ben Salem"


def test_extraire_linkedin():
    url = extraire_linkedin(CV_TEXTE_FR)
    assert url == "https://www.linkedin.com/in/amira-bensalem"


def test_extraire_linkedin_formats_varies():
    """Formats réels : sous-domaine régional, query, sans '/in/', libellé seul."""
    attendu = "https://www.linkedin.com/in/john-smith"
    assert extraire_linkedin("fr.linkedin.com/in/john-smith") == attendu
    assert extraire_linkedin("linkedin.com/in/john-smith?originalSubdomain=fr") == attendu
    assert extraire_linkedin("linkedin.com/john-smith") == attendu   # sans /in/
    assert extraire_linkedin("LinkedIn : john-smith") == attendu     # libellé + handle
    assert extraire_linkedin("LinkedIn: @john-smith") == attendu     # arobase


def test_extraire_linkedin_pas_de_faux_positif():
    """Un nom isolé après « LinkedIn » (sans handle) ne doit rien produire."""
    assert extraire_linkedin("Profil LinkedIn : John Smith") is None
    assert extraire_linkedin("Retrouvez-moi sur LinkedIn") is None


def test_extraire_localisation_ligne_seule():
    """La ville ne doit pas déborder sur la ligne précédente (bug corrigé)."""
    assert extraire_localisation(CV_TEXTE_FR) == "Ariana, Tunisie"


def test_extraire_identite_champs_absents_none():
    """Convention J10 : champ introuvable = None (pas de crash, pas de '')."""
    identite = extraire_identite("Texte sans aucune info personnelle.")
    assert identite.email is None
    assert identite.phone is None
    assert identite.linkedin is None


# ---------- langue + parser complet ----------
def test_detecter_langue_fr_en():
    assert detecter_langue(CV_TEXTE_FR) == "fr"
    assert detecter_langue(
        "Senior software engineer with ten years of experience building web services."
    ) == "en"


def test_detecter_langue_texte_trop_court():
    assert detecter_langue("abc") is None


def test_parser_cv_txt_bout_en_bout():
    cv = parser_cv("cv_amira.txt", CV_TEXTE_FR.encode("utf-8"))
    assert cv.id.startswith("cv_")
    assert cv.format == "txt"
    assert cv.language == "fr"
    assert cv.identity.first_name == "Amira"
    assert cv.identity.email == "amira.bensalem@gmail.com"
    assert cv.identity.phone == "+216 22 345 678"
    # nouveau en S4 : la section expérience est structurée en items
    assert len(cv.experience) == 1
    assert cv.experience[0].title == "Stage Data Analyst"
    assert cv.experience[0].company == "Vermeg"
    assert cv.experience[0].start_date == "2025-06"
    assert cv.experience[0].end_date == "2025-08"
