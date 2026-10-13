# -*- coding: utf-8 -*-
"""Tests S3 des endpoints FastAPI avec TestClient — pas besoin de lancer uvicorn."""
import io

import pytest
from fastapi.testclient import TestClient
from pypdf import PdfWriter

from app.main import app

client = TestClient(app)

CV_TEXTE = (
    "AMIRA BEN SALEM\n"
    "Ariana, Tunisie\n"
    "amira.bensalem@gmail.com\n"
    "+216 22 345 678\n"
)


@pytest.fixture
def pdf_vide() -> bytes:
    """Génère en mémoire un vrai PDF (1 page blanche) pour tester l'upload."""
    writer = PdfWriter()
    writer.add_blank_page(width=595, height=842)  # A4
    buf = io.BytesIO()
    writer.write(buf)
    return buf.getvalue()


@pytest.fixture
def pdf_avec_texte() -> bytes:
    """PDF réel contenant un mini-CV (généré avec fpdf2)."""
    from fpdf import FPDF
    pdf = FPDF()
    pdf.add_page()
    pdf.set_font("helvetica", size=11)
    for ligne in CV_TEXTE.split("\n"):
        pdf.cell(0, 7, ligne, new_x="LMARGIN", new_y="NEXT")
    return bytes(pdf.output())


@pytest.fixture
def docx_cv() -> bytes:
    """DOCX réel contenant le même mini-CV (généré avec python-docx)."""
    from docx import Document
    doc = Document()
    for ligne in CV_TEXTE.split("\n"):
        doc.add_paragraph(ligne)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_health():
    r = client.get("/health")
    assert r.status_code == 200
    assert r.json()["status"] == "ok"


def test_parse_pdf_vide_structure_cvdata(pdf_vide):
    """PDF blanc -> aucune info, mais la structure CVData doit être complète."""
    r = client.post("/api/cv/parse", files={"file": ("test.pdf", pdf_vide, "application/pdf")})
    assert r.status_code == 200
    data = r.json()
    assert data["id"].startswith("cv_")
    assert data["format"] == "pdf"
    for cle in ("identity", "experience", "education", "skills", "languages"):
        assert cle in data
    assert data["identity"]["email"] is None  # convention J10 : introuvable = None


def test_parse_pdf_extraction_identite(pdf_avec_texte):
    """Le cœur de la S3 : nom, email, téléphone extraits d'un vrai PDF."""
    r = client.post(
        "/api/cv/parse", files={"file": ("amira.pdf", pdf_avec_texte, "application/pdf")}
    )
    assert r.status_code == 200
    identite = r.json()["identity"]
    assert identite["first_name"] == "Amira"
    assert identite["last_name"] == "Ben Salem"
    assert identite["email"] == "amira.bensalem@gmail.com"
    assert identite["phone"] == "+216 22 345 678"


def test_parse_docx(docx_cv):
    """Nouveau en S3 : le format DOCX est accepté."""
    r = client.post(
        "/api/cv/parse",
        files={
            "file": (
                "amira.docx",
                docx_cv,
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
        },
    )
    assert r.status_code == 200
    data = r.json()
    assert data["format"] == "docx"
    assert data["identity"]["email"] == "amira.bensalem@gmail.com"


def test_parse_txt():
    """Nouveau en S3 : le format TXT est accepté."""
    r = client.post(
        "/api/cv/parse", files={"file": ("cv.txt", CV_TEXTE.encode("utf-8"), "text/plain")}
    )
    assert r.status_code == 200
    data = r.json()
    assert data["format"] == "txt"
    assert data["identity"]["phone"] == "+216 22 345 678"


def test_parse_mauvais_format():
    r = client.post("/api/cv/parse", files={"file": ("photo.png", b"fake", "image/png")})
    assert r.status_code == 400


def test_parse_pdf_corrompu():
    r = client.post("/api/cv/parse", files={"file": ("faux.pdf", b"pas un pdf", "application/pdf")})
    assert r.status_code == 422


def test_parse_sans_fichier():
    r = client.post("/api/cv/parse")
    assert r.status_code == 422  # validation FastAPI


def test_get_cv_apres_parse(pdf_avec_texte):
    r = client.post(
        "/api/cv/parse", files={"file": ("amira.pdf", pdf_avec_texte, "application/pdf")}
    )
    cv_id = r.json()["id"]
    r2 = client.get(f"/api/cv/{cv_id}")
    assert r2.status_code == 200
    assert r2.json()["id"] == cv_id
    assert r2.json()["identity"]["first_name"] == "Amira"


def test_get_cv_inconnu():
    r = client.get("/api/cv/cv_inexistant")
    assert r.status_code == 404


def test_batch_formats_melanges(pdf_avec_texte, docx_cv):
    """Un batch peut mélanger PDF, DOCX et TXT."""
    files = [
        ("files", ("a.pdf", pdf_avec_texte, "application/pdf")),
        ("files", ("b.docx", docx_cv, "application/octet-stream")),
        ("files", ("c.txt", CV_TEXTE.encode("utf-8"), "text/plain")),
    ]
    r = client.post("/api/cv/batch", files=files)
    assert r.status_code == 200
    data = r.json()
    assert len(data) == 3
    assert [cv["format"] for cv in data] == ["pdf", "docx", "txt"]


def test_pipeline_robuste_sans_dependances(pdf_vide):
    """Le pipeline S2 doit toujours répondre 200 même sans chromadb / clé Groq."""
    r = client.post("/api/pipeline", files={"file": ("test.pdf", pdf_vide, "application/pdf")})
    assert r.status_code == 200
    data = r.json()
    assert "entites" in data
    assert isinstance(data["avertissements"], list)
